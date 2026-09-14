from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = "/Users/yan/Desktop/gateways_dev_frontend/deliverables/Gateways_2026_Live_Demo_Script.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_border(cell, color="D9D9D9", sz="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        el = borders.find(tag)
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def add_field(paragraph, field):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    separate.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)


def add_labelled_line(doc, label, text):
    p = doc.add_paragraph(style="Spoken")
    r = p.add_run(label + " ")
    r.bold = True
    p.add_run(text)
    return p


def add_action(doc, text):
    p = doc.add_paragraph(style="Action")
    p.add_run("On screen: ").bold = True
    p.add_run(text)
    return p


def add_note(doc, text):
    p = doc.add_paragraph(style="Note")
    p.add_run("Demo note: ").bold = True
    p.add_run(text)
    return p


def add_step(doc, heading, actions):
    doc.add_heading(heading, level=1)
    for kind, text in actions:
        if kind == "action":
            add_action(doc, text)
        elif kind == "say":
            add_labelled_line(doc, "Say:", f'“{text}”')
        elif kind == "note":
            add_note(doc, text)
        elif kind == "bullet":
            doc.add_paragraph(text, style="List Bullet")


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.68)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(5)

    for style_name, size, before, after in (("Heading 1", 15, 15, 5), ("Heading 2", 12, 10, 3)):
        s = styles[style_name]
        s.font.name = "Aptos Display"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(before)
        s.paragraph_format.space_after = Pt(after)
        s.paragraph_format.keep_with_next = True

    action = styles.add_style("Action", WD_STYLE_TYPE.PARAGRAPH)
    action.base_style = styles["Normal"]
    action.font.name = "Aptos"
    action.font.size = Pt(9.5)
    action.font.italic = True
    action.font.color.rgb = RGBColor(89, 89, 89)
    action.paragraph_format.left_indent = Inches(0.22)
    action.paragraph_format.first_line_indent = Inches(-0.22)
    action.paragraph_format.space_before = Pt(5)
    action.paragraph_format.space_after = Pt(2)

    spoken = styles.add_style("Spoken", WD_STYLE_TYPE.PARAGRAPH)
    spoken.base_style = styles["Normal"]
    spoken.font.size = Pt(10.5)
    spoken.paragraph_format.left_indent = Inches(0.22)
    spoken.paragraph_format.first_line_indent = Inches(-0.22)
    spoken.paragraph_format.space_after = Pt(5)

    note = styles.add_style("Note", WD_STYLE_TYPE.PARAGRAPH)
    note.base_style = styles["Normal"]
    note.font.size = Pt(9.5)
    note.font.color.rgb = RGBColor(89, 89, 89)
    note.paragraph_format.space_before = Pt(5)
    note.paragraph_format.space_after = Pt(4)

    # Header and footer
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("GATEWAYS 2026   |   PARALLAX")
    hr.font.name = "Aptos"
    hr.font.size = Pt(8)
    hr.font.bold = True
    hr.font.color.rgb = RGBColor(89, 89, 89)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("Live Demo Script   •   Page ")
    fr.font.name = "Aptos"
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor(89, 89, 89)
    add_field(footer, "PAGE")

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("Gateways 2026 Live Demo Script")
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(10)
    r = sub.add_run("Presenter guide for the Gateways 2026 Parallax website")
    r.italic = True
    r.font.size = Pt(12)

    intro = doc.add_paragraph()
    intro.add_run("Purpose. ").bold = True
    intro.add_run("Use this script to take the audience from the first loading screen through the public site, account journey, participant dashboard, team features, and 3D campus. Read the spoken lines naturally and keep the on-screen actions moving.")

    doc.add_heading("Before You Start", level=1)
    prep = doc.add_table(rows=1, cols=2)
    prep.alignment = WD_TABLE_ALIGNMENT.CENTER
    prep.style = "Table Grid"
    prep.autofit = False
    prep.columns[0].width = Inches(2.0)
    prep.columns[1].width = Inches(5.0)
    hdr = prep.rows[0].cells
    for cell, text in zip(hdr, ("Prepare", "What to have ready")):
        cell.width = Inches(2.0) if cell is hdr[0] else Inches(5.0)
        set_cell_shading(cell, "33245A")
        set_cell_margins(cell)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(text)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9.5)
    set_repeat_table_header(prep.rows[0])
    prep_rows = [
        ("Fresh browser tab", "Use a fresh tab so the opening animation appears. Keep the site on desktop view for the top navigation and 3D world."),
        ("Demo participant", "Use a prepared account with completed participant details and a verified payment status."),
        ("Team example", "Have one team event registration ready with a visible team name and invite code."),
        ("Organiser example", "Optional: prepare a separate organiser account and resettable receipt for the payment-verification walkthrough."),
        ("Safe demonstration", "Do not submit a real payment, remove a real teammate, disband a real team, or approve or reject a real receipt during the demo."),
    ]
    for idx, (left, right) in enumerate(prep_rows):
        cells = prep.add_row().cells
        for i, (cell, text) in enumerate(zip(cells, (left, right))):
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx % 2 == 1:
                set_cell_shading(cell, "F3F5F8")
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.size = Pt(9.5)
            if i == 0:
                run.bold = True

    add_step(doc, "1 Opening Screen", [
        ("action", "Open the website in a fresh browser tab. Let the crest animation play for a moment; click once only if you need to skip it."),
        ("say", "Welcome to Gateways 2026, Parallax. The experience begins with the fest identity assembling on screen before opening into the Gateways world."),
        ("action", "Pause briefly on the homepage hero. Point to the announcement ticker, date, countdown, and Parallax title."),
        ("say", "The homepage immediately introduces the fest dates, live announcements, countdown, and our theme for the year, Digital Twins."),
    ])

    add_step(doc, "2 Homepage Controls and Public Information", [
        ("action", "Click the music control once, then click the appearance toggle to demonstrate the alternate colour theme."),
        ("say", "Visitors can control the background music and switch between light and dark appearances at any time."),
        ("action", "Open Events from the navigation. Click All, Technical, and Non Technical filters. Open one event card, then return to the event list."),
        ("say", "The event browser lets visitors filter the lineup and open any event for its details, date, participation format, rules, prizes, and registration path."),
        ("action", "Open Schedule from the navigation and scroll through both dates. Close the modal."),
        ("say", "The schedule is available directly from the homepage and groups every event by date."),
        ("action", "Scroll through the fest overview, Digital Twins section, entry-fee information, registration steps, and contacts."),
        ("say", "The site explains what Gateways is, this year’s theme, the entry fee, accommodation, the registration process, and the right contact for every type of question."),
        ("action", "Open Gallery. Change editions and use previous and next controls. Then open FAQ and expand one question."),
        ("say", "The gallery keeps memories from past editions accessible, while the FAQ gives visitors quick answers without leaving the site."),
        ("action", "Open Contact, then point to the footer links for guidelines, registration process, policies, and social channels."),
        ("say", "Visitors can reach the organising team directly, or use the footer to find the supporting information they need."),
    ])

    add_step(doc, "3 Pixey Assistant", [
        ("action", "Click Ask Pixey. Click one quick question such as What are the fees, then type a short question. Show Clear and close the assistant."),
        ("say", "Pixey is the built in Gateways assistant. It answers common questions about events, fees, dates, and registration using the information on this site."),
        ("say", "A visitor can choose a quick question, type their own question, clear the conversation, or close the assistant whenever they wish."),
    ])

    add_step(doc, "4 Account Journey", [
        ("action", "Click Start the Journey and let the portal transition appear. On the login page, open Sign Up."),
        ("say", "When a visitor is ready to take part, Start the Journey takes them through the portal and into the account area."),
        ("action", "Point to email, character name, password checklist, confirm-password feedback, and show-password controls. Do not create a real account."),
        ("say", "New visitors can create an account with an email, character name, and password. The form gives immediate guidance while they complete it."),
        ("action", "Switch to Login. Point to Remember me and Forgot password. Use the prepared demo account to sign in."),
        ("say", "Returning visitors can sign in, choose to be remembered on the device, or start the password-recovery flow if required."),
    ])

    add_step(doc, "5 Participant Profile and Payment", [
        ("action", "After sign-in, let the transition finish and open Profile. Show player identity, progress, attendance, and profile status."),
        ("say", "This is the participant profile. It brings together the player identity, progress, event attendance, and registration readiness."),
        ("action", "Open Complete your details or Edit details. Briefly show the form, then close it without changing data."),
        ("say", "Participant details are collected once and then reused for every event registration."),
        ("action", "Open Submit Payment Details. Show the instructions, video tutorial option, transaction ID field, receipt upload, and optional screenshot upload. Close it without submitting."),
        ("say", "The payment area gives clear instructions, a tutorial, and a simple place to submit the transaction details and receipt. After submission, the participant can see when the receipt is under review."),
        ("note", "Use only test files and resettable data if you need to demonstrate a completed payment submission."),
    ])

    add_step(doc, "6 Explore Events and Register", [
        ("action", "Open Explore Events. Search for an event by name or type, clear the search, and use each event-track filter."),
        ("say", "Participants can search the event directory and filter it to find exactly what they want to join."),
        ("action", "Open an event card. Point to its description, date, venue, rules link, and registration control."),
        ("say", "Each event page gives participants the information they need before committing to it."),
        ("action", "For an individual event, open the registration confirmation and cancel it unless you are using resettable test data."),
        ("say", "Individual events use a short confirmation step before registration is completed."),
        ("action", "For a team event, show Create Team and Join Team. Open the team-name field and the invite-code field. If test data is available, show the team success screen and copy or share the code."),
        ("say", "For team events, a participant can either create a team and receive an invite code, or join an existing team using that code."),
    ])

    add_step(doc, "7 My Events and Team Management", [
        ("action", "Open My Events. Point to Upcoming and Past events. Open Manage on a prepared team event."),
        ("say", "My Events keeps every registration in one place and separates upcoming activity from past activity."),
        ("action", "Show the team name, invite code, Copy Code, Share Invite, and roster. Point to leader management controls but do not make destructive changes."),
        ("say", "Team leaders can view the roster, share the invite code, and manage the team. Confirmation is required before changes such as leaving, removing a member, or disbanding a team."),
        ("action", "Open Join Team and show where an invite code is entered, then close it."),
        ("say", "Participants can also join a team later by entering a valid invite code."),
    ])

    add_step(doc, "8 Interactive Campus", [
        ("action", "Open 3D World. Allow the campus to load. Walk briefly, enter a classroom, and press E to open its event hub."),
        ("say", "The 3D campus turns every classroom into an event space. A participant can walk into a room and press E to open that event’s hub."),
        ("action", "In the event hub, point to event details, rules, and the registration control. Close it."),
        ("say", "The same event information and registration path is available inside the campus, so the experience stays consistent."),
        ("action", "Switch from 3D to List view. Open one event from the list, then close it. Click Inventory to return to the participant area."),
        ("say", "List view provides the same campus events in a quick browser format for participants who prefer not to walk through the world."),
    ])

    add_step(doc, "9 Dashboard Utilities", [
        ("action", "Open Schedule, Announcements, Achievements or Inventory if populated, Notifications, Guidelines, FAQ, Terms, and Privacy. Keep this section brisk."),
        ("say", "The remaining dashboard sections keep the schedule, announcements, notifications, progress, guidelines, and participant information readily available."),
        ("action", "Open Settings. Click Light, Dark, and Follow system. Then show Full motion, Reduce motion, and Follow system."),
        ("say", "Settings allow each participant to choose their preferred appearance and motion level, including following their device settings."),
        ("action", "Return to the sidebar and sign out only as the final action."),
        ("say", "When the participant is finished, they can sign out securely from the dashboard."),
    ])

    add_step(doc, "10 Organiser Payment Verification Optional", [
        ("action", "Sign in with the prepared organiser account. Open payment verification and choose a resettable pending receipt."),
        ("say", "Organisers have a separate payment-verification area for reviewing submitted receipts."),
        ("action", "Show the receipt view, Approve, and Reject. Open the rejection-reason form, then close it without saving unless the record is resettable."),
        ("say", "An organiser can approve a valid receipt or return it with a clear reason. Once approved, the participant can complete event registration."),
    ])

    doc.add_heading("Final Checks and Fallbacks", level=1)
    add_note(doc, "The Read Brochure and Google sign in controls are marked as coming soon or disabled. Do not present them as working live-demo features.")
    doc.add_paragraph("If an interaction does not load during the demo, keep the audience moving with the nearest working alternative.")
    for item in [
        "If the opening animation is too slow, click once to move directly to the homepage.",
        "If sign in is unavailable, continue with the public homepage, event browser, schedule, gallery, FAQ, and Pixey assistant.",
        "If the 3D campus is not available on the browser, switch to List view and open an event hub from there.",
        "If an event list is still loading, show the schedule or public information while it finishes.",
        "If a test account does not show payment verification or team controls, describe the flow and move on rather than altering real participant data.",
        "Keep the organiser payment-verification section optional unless a resettable test receipt is ready.",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Closing Line", level=1)
    add_labelled_line(doc, "Say:", "That is Gateways 2026 Parallax: a complete fest journey that starts with discovery, guides participants through registration, and gives them a personalised place to explore events, teams, updates, and the campus.")

    doc.save(OUT)


if __name__ == "__main__":
    build()
